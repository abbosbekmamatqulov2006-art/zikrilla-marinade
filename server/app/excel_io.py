"""Excel orqali ommaviy tahrirlash: mahsulotlar, kategoriyalar, filiallar (3 tilda).

Eksport: joriy ma'lumotlar bilan to'ldirilgan shablon (.xlsx).
Import: o'sha shablonni o'qiydi. Avval tekshiradi (apply=False), keyin qo'llaydi (apply=True).
Qoidalar:
  - Rasm va videolar Excel'da yo'q — ular panelda qo'lda yuklanadi va import ularga tegmaydi.
  - ID bo'sh qator — yangi yozuv. ID bor — o'sha yozuv yangilanadi.
  - Excel'da yo'q mahsulot/filial o'chirilmaydi. Yashirish uchun "Faol" = yo'q.
  - Bitta xato bo'lsa ham hech narsa yozilmaydi (hammasi yoki hech narsa).
  - RU/EN katakchasi bo'sh bo'lsa, UZ matni qo'yiladi.
"""
import datetime as dt
import io
import re
import uuid

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from .models import Product, Store
from .site_settings import get_site, save_site

LANGS = ("uz", "ru", "en")
L = {"uz": "UZ", "ru": "RU", "en": "EN"}

P_COLS = (["ID", "Faol", "Tartib", "Kategoriya", "Narx (so'm)"]
          + [f"Nomi {L[x]}" for x in LANGS]
          + [f"Birlik {L[x]}" for x in LANGS]
          + [f"Qisqa tavsif {L[x]}" for x in LANGS]
          + [f"Tarkibi {L[x]}" for x in LANGS]
          + [f"Batafsil tavsif {L[x]}" for x in LANGS]
          + ["Achchiq", "Haftalik aksiya", "Kam qolgan", "Pishirish (daqiqa)", "Rasm/video (panelda)"])
C_COLS = ["ID"] + [f"Nomi {L[x]}" for x in LANGS]
S_COLS = (["ID", "Faol", "Tartib"]
          + [f"Nomi {L[x]}" for x in LANGS]
          + [f"Manzil {L[x]}" for x in LANGS]
          + ["Kenglik (lat)", "Uzunlik (lng)", "Telefon", "Ochilish", "Yopilish"])

HEAD_FILL = PatternFill("solid", fgColor="D50C05")
HEAD_FONT = Font(bold=True, color="FFFFFF")
RO_FILL = PatternFill("solid", fgColor="EEEEEE")

GUIDE = [
    "ZIKRILLA MARINADE — Excel orqali tahrirlash",
    "",
    "1. Shu faylni paneldan yuklab oling (unda hozirgi ma'lumotlar bor).",
    "2. Kerakli varaqni tahrirlang: Mahsulotlar, Kategoriyalar, Filiallar.",
    "3. Panel → Mahsulotlar → «Excel yuklash». Avval o'zgarishlar ro'yxati chiqadi, «Tasdiqlash» bosilgandagina yoziladi.",
    "",
    "QOIDALAR",
    "• Yangi mahsulot/filial: ID katakchasini bo'sh qoldiring — ID avtomatik beriladi.",
    "• Mavjudini o'zgartirish: ID ga tegmang, boshqa katakchalarni o'zgartiring.",
    "• O'chirish o'rniga yashirish: «Faol» = yo'q. Excel'dan qatorni o'chirsangiz, mahsulot saytda qolaveradi.",
    "• Rasm va videolar Excel'da yo'q — ularni panelda mahsulotni ochib qo'lda yuklang. Import ularga tegmaydi.",
    "• RU yoki EN bo'sh bo'lsa, UZ matni qo'yiladi.",
    "• Kategoriya ustuniga Kategoriyalar varag'idagi ID yoziladi (masalan: skewers).",
    "• Narx — faqat raqam, so'mda (masalan: 145000).",
    "• Ha/yo'q ustunlari: ha yoki yo'q.",
    "• Filial koordinatasi: Google xaritada nuqtani bosing → 41.3111, 69.2797 ko'rinishidagi sonlar.",
    "• Ish vaqti: 09:00 ko'rinishida. 24 soat ishlasa: 00:00 va 23:59.",
    "• Kategoriyalar varag'idagi tartib — saytdagi tartib. Mahsulot ishlatayotgan kategoriyani o'chirib bo'lmaydi.",
    "• Biror qatorda xato bo'lsa, hech narsa yozilmaydi va xato qaysi qatorda ekani ko'rsatiladi.",
    "• Har importdan oldin server avvalgi holatning nusxasini saqlaydi.",
]


# ------------------------------------------------------------------ eksport
def _sheet(wb, title, cols, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 32
    for i, name in enumerate(cols, start=1):
        w = (widths or {}).get(name)
        if w is None:
            w = 40 if "tavsif" in name.lower() or "Tarkibi" in name or "Manzil" in name else 18
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def _yes(v: bool) -> str:
    return "ha" if v else "yo'q"


def _yesno_validation(ws, cols, names, last_row=500):
    dv = DataValidation(type="list", formula1='"ha,yo\'q"', allow_blank=True)
    ws.add_data_validation(dv)
    for n in names:
        col = get_column_letter(cols.index(n) + 1)
        dv.add(f"{col}2:{col}{last_row}")


def export_xlsx(db: Session) -> bytes:
    site = get_site(db)
    cats = site.get("categories") or []
    wb = Workbook()
    g = wb.active
    g.title = "Yo'riqnoma"
    for line in GUIDE:
        g.append([line])
    g["A1"].font = Font(bold=True, size=14)
    g["A7"].font = Font(bold=True)
    g.column_dimensions["A"].width = 120

    prows = []
    for p in db.query(Product).order_by(Product.sort, Product.id).all():
        def lx(d, k):
            return (d or {}).get(k, "")
        prows.append([p.id, _yes(p.active), p.sort or 0, p.category, p.price]
                     + [lx(p.name, x) for x in LANGS]
                     + [lx(p.unit, x) for x in LANGS]
                     + [lx(p.subtitle, x) for x in LANGS]
                     + [lx(p.ingredients, x) for x in LANGS]
                     + [lx(p.description, x) for x in LANGS]
                     + [_yes(p.spicy), _yes(p.weekly_deal), _yes(p.low_stock), p.grill_minutes,
                        len(p.media or []) or (1 if p.image else 0)])
    ws = _sheet(wb, "Mahsulotlar", P_COLS, prows, {"ID": 22, "Faol": 8, "Tartib": 8, "Narx (so'm)": 13})
    _yesno_validation(ws, P_COLS, ["Faol", "Achchiq", "Haftalik aksiya", "Kam qolgan"])
    if cats:
        dv = DataValidation(type="list", formula1=f"=Kategoriyalar!$A$2:$A${len(cats) + 40}", allow_blank=False)
        ws.add_data_validation(dv)
        col = get_column_letter(P_COLS.index("Kategoriya") + 1)
        dv.add(f"{col}2:{col}500")
    ro = P_COLS.index("Rasm/video (panelda)") + 1
    for r in range(2, len(prows) + 2):
        ws.cell(row=r, column=ro).fill = RO_FILL

    _sheet(wb, "Kategoriyalar", C_COLS,
           [[c.get("id", "")] + [(c.get("name") or {}).get(x, "") for x in LANGS] for c in cats],
           {"ID": 18})

    srows = []
    for s in db.query(Store).order_by(Store.sort, Store.id).all():
        srows.append([s.id, _yes(s.active), s.sort or 0]
                     + [(s.name or {}).get(x, "") for x in LANGS]
                     + [(s.address or {}).get(x, "") for x in LANGS]
                     + [s.lat, s.lng, s.phone or "", s.open, s.close])
    ws = _sheet(wb, "Filiallar", S_COLS, srows, {"ID": 18, "Faol": 8, "Tartib": 8})
    _yesno_validation(ws, S_COLS, ["Faol"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ import
class RowError(Exception):
    pass


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _bool(v, default=False) -> bool:
    t = _s(v).lower().replace("‘", "'").replace("’", "'")
    if t == "":
        return default
    if t in ("ha", "1", "true", "yes", "да", "+", "x"):
        return True
    if t in ("yo'q", "yoq", "0", "false", "no", "нет", "-"):
        return False
    raise RowError(f"«{_s(v)}» — ha yoki yo'q bo'lishi kerak")


def _int(v, field, default=None, lo=None) -> int:
    t = _s(v).replace(" ", "").replace(" ", "").replace(",", "")
    if t == "":
        if default is None:
            raise RowError(f"{field} bo'sh")
        return default
    try:
        n = int(float(t))
    except ValueError:
        raise RowError(f"{field}: «{_s(v)}» raqam emas")
    if lo is not None and n < lo:
        raise RowError(f"{field} {lo} dan kichik bo'lmaydi")
    return n


def _float(v, field) -> float:
    t = _s(v).replace(",", ".")
    try:
        return float(t)
    except ValueError:
        raise RowError(f"{field}: «{_s(v)}» son emas")


def _time(v, field) -> str:
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, dt.datetime):
        return v.strftime("%H:%M")
    if isinstance(v, (int, float)) and 0 <= v < 1:      # Excel vaqt formati (kun ulushi)
        m = round(v * 24 * 60)
        return f"{m // 60:02d}:{m % 60:02d}"
    t = _s(v)
    m = re.match(r"^(\d{1,2})[:.](\d{2})", t)
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        raise RowError(f"{field}: «{t}» 09:00 ko'rinishida bo'lsin")
    return f"{int(m.group(1)):02d}:{m.group(2)}"


def _lx(row: dict, prefix: str) -> dict:
    d = {x: _s(row.get(f"{prefix} {L[x]}")) for x in LANGS}
    base = d["uz"] or d["ru"] or d["en"]
    return {k: (v or base) for k, v in d.items()}


def _rows(ws, cols):
    head = [_s(c.value) for c in ws[1]]
    missing = [c for c in cols if c not in head and c != "Rasm/video (panelda)"]
    if missing:
        raise RowError(f"«{ws.title}» varag'ida ustunlar yo'q: {', '.join(missing)}. Paneldan yangi shablonni yuklab oling.")
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = {head[j]: r[j] for j in range(min(len(head), len(r))) if head[j]}
        if all(_s(v) == "" for v in row.values()):
            continue
        yield i, row


def _norm(v):
    """None va bo'sh {uz:'',ru:'',en:''} bir xil hisoblanadi."""
    if isinstance(v, dict) and not any((x or "") for x in v.values()):
        return None
    return v


def _same(obj, field, value) -> bool:
    return _norm(getattr(obj, field)) == _norm(value)


def import_xlsx(db: Session, data: bytes, apply: bool) -> dict:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        return {"ok": False, "applied": False, "errors": [{"sheet": "", "row": 0, "message": "Fayl o'qilmadi. .xlsx shablonidan foydalaning."}]}

    errors: list[dict] = []
    report = {s: {"created": [], "updated": [], "unchanged": 0} for s in ("Mahsulotlar", "Kategoriyalar", "Filiallar")}

    def err(sheet, row, msg):
        errors.append({"sheet": sheet, "row": row, "message": msg})

    site = get_site(db)
    new_cats = None

    # ---- kategoriyalar
    if "Kategoriyalar" in wb.sheetnames:
        try:
            cats, seen = [], set()
            for i, row in _rows(wb["Kategoriyalar"], C_COLS):
                cid = _s(row.get("ID")).lower()
                name = _lx(row, "Nomi")
                if not name["uz"]:
                    err("Kategoriyalar", i, "Nomi bo'sh"); continue
                if not cid:
                    cid = re.sub(r"[^a-z0-9]+", "-", name["uz"].lower()).strip("-")[:24] or f"cat-{uuid.uuid4().hex[:4]}"
                if not re.match(r"^[a-z0-9_-]{1,32}$", cid):
                    err("Kategoriyalar", i, f"ID «{cid}»: faqat lotin harflari, raqam va - _"); continue
                if cid in seen:
                    err("Kategoriyalar", i, f"ID «{cid}» ikki marta yozilgan"); continue
                seen.add(cid)
                cats.append({"id": cid, "name": name})
            if cats:
                old = {c["id"]: c for c in site.get("categories") or []}
                for c in cats:
                    if c["id"] not in old:
                        report["Kategoriyalar"]["created"].append(c["id"])
                    elif old[c["id"]].get("name") != c["name"]:
                        report["Kategoriyalar"]["updated"].append(c["id"])
                    else:
                        report["Kategoriyalar"]["unchanged"] += 1
                removed = set(old) - seen
                for rid in removed:
                    if db.query(Product).filter(Product.category == rid).first():
                        err("Kategoriyalar", 0, f"«{rid}» kategoriyasi o'chirilgan, lekin unda mahsulotlar bor. "
                                                 "Avval mahsulotlarni boshqa kategoriyaga o'tkazing.")
                report["Kategoriyalar"]["removed"] = sorted(removed)
                new_cats = cats
        except RowError as e:
            err("Kategoriyalar", 0, str(e))
    cat_ids = {c["id"] for c in (new_cats if new_cats is not None else site.get("categories") or [])}

    # ---- mahsulotlar
    if "Mahsulotlar" in wb.sheetnames:
        try:
            seen = set()
            for i, row in _rows(wb["Mahsulotlar"], P_COLS):
                try:
                    pid = _s(row.get("ID"))
                    name = _lx(row, "Nomi")
                    if not name["uz"]:
                        raise RowError("Nomi bo'sh")
                    cat = _s(row.get("Kategoriya")).lower()
                    if cat not in cat_ids:
                        raise RowError(f"Kategoriya «{cat}» yo'q. Kategoriyalar varag'idagi ID ni yozing.")
                    vals = {
                        "active": _bool(row.get("Faol"), True),
                        "sort": _int(row.get("Tartib"), "Tartib", 0),
                        "category": cat,
                        "price": _int(row.get("Narx (so'm)"), "Narx", lo=0),
                        "name": name,
                        "subtitle": _lx(row, "Qisqa tavsif"),
                        "ingredients": _lx(row, "Tarkibi"),
                        "description": _lx(row, "Batafsil tavsif"),
                        "spicy": _bool(row.get("Achchiq")),
                        "weekly_deal": _bool(row.get("Haftalik aksiya")),
                        "low_stock": _bool(row.get("Kam qolgan")),
                        "grill_minutes": _int(row.get("Pishirish (daqiqa)"), "Pishirish", 12, lo=0),
                    }
                    unit = _lx(row, "Birlik")
                    vals["unit"] = unit if unit["uz"] else None
                    vals["is_extra"] = cat == "extras"
                    p = db.get(Product, pid) if pid else None
                    if pid and pid in seen:
                        raise RowError(f"ID «{pid}» ikki marta yozilgan")
                    if p is None:
                        if pid and not re.match(r"^[a-z0-9-]{1,32}$", pid):
                            raise RowError(f"ID «{pid}» topilmadi. Yangi mahsulot uchun ID ni bo'sh qoldiring.")
                        if not pid:
                            pid = re.sub(r"[^a-z0-9]+", "-", name["uz"].lower()).strip("-")[:24] or "p"
                            if db.get(Product, pid) is not None or pid in seen:
                                pid = f"{pid}-{uuid.uuid4().hex[:5]}"
                        p = Product(id=pid, rating=4.8, reviews=0, image="", media=[])
                        db.add(p)
                        report["Mahsulotlar"]["created"].append(f"{pid} — {name['uz']}")
                        changed = True
                    else:
                        changed = any(not _same(p, k, v) for k, v in vals.items())
                        if changed:
                            report["Mahsulotlar"]["updated"].append(f"{pid} — {name['uz']}")
                        else:
                            report["Mahsulotlar"]["unchanged"] += 1
                    seen.add(pid)
                    if changed:
                        for k, v in vals.items():
                            setattr(p, k, v)
                except RowError as e:
                    err("Mahsulotlar", i, str(e))
        except RowError as e:
            err("Mahsulotlar", 0, str(e))

    # ---- filiallar
    if "Filiallar" in wb.sheetnames:
        try:
            seen = set()
            for i, row in _rows(wb["Filiallar"], S_COLS):
                try:
                    sid = _s(row.get("ID"))
                    name = _lx(row, "Nomi")
                    if not name["uz"]:
                        raise RowError("Nomi bo'sh")
                    lat = _float(row.get("Kenglik (lat)"), "Kenglik")
                    lng = _float(row.get("Uzunlik (lng)"), "Uzunlik")
                    if not (-90 <= lat <= 90 and -180 <= lng <= 180):
                        raise RowError("Koordinatalar noto'g'ri")
                    phone = re.sub(r"[^\d+]", "", _s(row.get("Telefon")))[:16]
                    vals = {
                        "active": _bool(row.get("Faol"), True),
                        "sort": _int(row.get("Tartib"), "Tartib", 0),
                        "name": name, "address": _lx(row, "Manzil"),
                        "lat": lat, "lng": lng, "phone": phone,
                        "open": _time(row.get("Ochilish") or "09:00", "Ochilish"),
                        "close": _time(row.get("Yopilish") or "22:00", "Yopilish"),
                    }
                    s = db.get(Store, sid) if sid else None
                    if sid and sid in seen:
                        raise RowError(f"ID «{sid}» ikki marta yozilgan")
                    if s is None:
                        if sid and not re.match(r"^[a-z0-9]{1,32}$", sid):
                            raise RowError(f"ID «{sid}» topilmadi. Yangi filial uchun ID ni bo'sh qoldiring.")
                        if not sid:
                            sid = re.sub(r"[^a-z0-9]+", "", name["uz"].lower())[:20] or "filial"
                            if db.get(Store, sid) is not None or sid in seen:
                                sid = f"{sid}{uuid.uuid4().hex[:4]}"
                        s = Store(id=sid)
                        db.add(s)
                        report["Filiallar"]["created"].append(f"{sid} — {name['uz']}")
                        changed = True
                    else:
                        changed = any(not _same(s, k, v) for k, v in vals.items())
                        if changed:
                            report["Filiallar"]["updated"].append(f"{sid} — {name['uz']}")
                        else:
                            report["Filiallar"]["unchanged"] += 1
                    seen.add(sid)
                    if changed:
                        for k, v in vals.items():
                            setattr(s, k, v)
                except RowError as e:
                    err("Filiallar", i, str(e))
        except RowError as e:
            err("Filiallar", 0, str(e))

    ok = not errors
    if ok and apply:
        if new_cats is not None:
            site["categories"] = new_cats
            save_site(db, site)          # commit — mahsulot va filiallar ham shu bilan yoziladi
        db.commit()
    else:
        db.rollback()
    return {"ok": ok, "applied": ok and apply, "errors": errors[:200], "report": report}
